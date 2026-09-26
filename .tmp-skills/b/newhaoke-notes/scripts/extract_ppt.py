#!/usr/bin/env python3
"""
PPT 路径预判与文字提取
判断 PPTX 是带文字版还是纯图版：
- 带文字版（python-pptx 直接提取文字字符数 > 阈值）→ 直接生成 ppt_md
- 纯图版（文字字符数 < 阈值）→ 返回需要走 OCR 路径的标志

同时返回每页的文字字符数和图片形状数，便于细粒度判断。

用法：
    # 单文件
    python extract_ppt.py --input raw/xxx_PPT.pptx --output data/ppt_md/第1课_ppt.md

    # 检查模式（不输出文件，只报告）
    python extract_ppt.py --input raw/xxx_PPT.pptx --check
"""

import argparse
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# 判断阈值：每页平均文字字符数
TEXT_THRESHOLD_PER_SLIDE = 30   # 平均每页 < 30 字 → 视为纯图


def analyze_pptx(pptx_path: str) -> dict:
    """
    分析 PPTX，返回每页的文字字符数和图片形状数。
    """
    prs = Presentation(pptx_path)
    pages = []

    for slide_idx, slide in enumerate(prs.slides, 1):
        text_chars = 0
        picture_count = 0
        text_blocks = []

        for shape in slide.shapes:
            # 提取文字
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    para_text = ''.join(run.text for run in para.runs).strip()
                    if para_text:
                        text_blocks.append(para_text)
                        text_chars += len(para_text)

            # 统计图片
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1

        pages.append({
            'slide_num': slide_idx,
            'text_chars': text_chars,
            'picture_count': picture_count,
            'text_blocks': text_blocks,
        })

    total_text = sum(p['text_chars'] for p in pages)
    total_pictures = sum(p['picture_count'] for p in pages)
    avg_text_per_slide = total_text / len(pages) if pages else 0

    is_text_pptx = avg_text_per_slide >= TEXT_THRESHOLD_PER_SLIDE

    return {
        'total_slides': len(pages),
        'total_text_chars': total_text,
        'total_pictures': total_pictures,
        'avg_text_per_slide': avg_text_per_slide,
        'is_text_pptx': is_text_pptx,
        'pages': pages,
    }


def render_text_to_md(analysis: dict) -> str:
    """将带文字 PPTX 的分析结果渲染为 Markdown"""
    lines = []
    for page in analysis['pages']:
        lines.append(f"## 幻灯片 {page['slide_num']}")
        if page['text_blocks']:
            for block in page['text_blocks']:
                lines.append(block)
        else:
            lines.append('（此页无文字内容）')
        lines.append('')
    return '\n'.join(lines)


def extract_to_md(pptx_path: str, output_path: str) -> dict:
    """
    分析 PPTX 并尝试直接提取文字。
    - is_text_pptx == True：直接写出 .md，返回 {'route': 'text', ...}
    - is_text_pptx == False：不写文件，返回 {'route': 'image', ...}，调用方应走 OCR 路径
    """
    analysis = analyze_pptx(pptx_path)

    if analysis['is_text_pptx']:
        md = render_text_to_md(analysis)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(md)

        return {
            'route': 'text',
            'output': output_path,
            'slides': analysis['total_slides'],
            'total_text_chars': analysis['total_text_chars'],
            'avg_text_per_slide': round(analysis['avg_text_per_slide'], 1),
        }
    else:
        return {
            'route': 'image',
            'output': None,
            'slides': analysis['total_slides'],
            'total_text_chars': analysis['total_text_chars'],
            'total_pictures': analysis['total_pictures'],
            'avg_text_per_slide': round(analysis['avg_text_per_slide'], 1),
            'reason': f"平均每页 {analysis['avg_text_per_slide']:.1f} 字 < {TEXT_THRESHOLD_PER_SLIDE}",
        }


def main():
    parser = argparse.ArgumentParser(description='PPT 路径预判与文字提取')
    parser.add_argument('--input', required=True, help='输入 PPTX 文件')
    parser.add_argument('--output', help='输出 .md 文件路径（仅 text 路径需要）')
    parser.add_argument('--check', action='store_true',
                        help='仅检查路径类型，不输出文件')
    args = parser.parse_args()

    if not os.path.isfile(args.input):
        print(f"[错误] 输入文件不存在: {args.input}")
        sys.exit(1)

    if args.input.lower().endswith('.ppt'):
        print(f"[错误] .ppt 格式不支持自动处理。")
        print(f"       请用 PowerPoint / WPS / Keynote 打开 {args.input}，")
        print(f"       手动另存为 .pptx 后再继续。不要尝试 libreoffice 等命令行工具。")
        sys.exit(1)

    if args.check:
        analysis = analyze_pptx(args.input)
        print(f"PPTX: {args.input}")
        print(f"  幻灯片数: {analysis['total_slides']}")
        print(f"  总文字字符: {analysis['total_text_chars']}")
        print(f"  总图片形状: {analysis['total_pictures']}")
        print(f"  平均每页文字: {analysis['avg_text_per_slide']:.1f}")
        print(f"  路径判断: {'带文字（python-pptx 提取）' if analysis['is_text_pptx'] else '纯图（需 GLM-OCR）'}")
        return

    if not args.output:
        print("[错误] 非 --check 模式必须提供 --output")
        sys.exit(1)

    result = extract_to_md(args.input, args.output)
    if result['route'] == 'text':
        print(f"  ✓ 带文字 PPT，已直接提取")
        print(f"    页数: {result['slides']}, 字符: {result['total_text_chars']}, "
              f"输出: {result['output']}")
    else:
        print(f"  → 纯图 PPT，需走 OCR 路径")
        print(f"    页数: {result['slides']}, 平均字符: {result['avg_text_per_slide']}")
        print(f"    {result['reason']}")
        # 用 exit code 2 表示需要 OCR
        sys.exit(2)


if __name__ == '__main__':
    main()
